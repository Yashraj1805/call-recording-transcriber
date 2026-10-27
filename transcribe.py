#!/usr/bin/env python3
"""
Fast parallel call-recording transcriber using Deepgram.

Transcribes every audio file (mp3/wav) in a folder using multiple parallel
workers instead of one-at-a-time, so a batch of thousands of recordings
finishes in minutes instead of hours. Writes results incrementally to a
JSONL file (safe to stop/resume) and can export a clean Excel report.

Usage:
    export DEEPGRAM_API_KEY=your_key_here      # or pass --api-key
    python transcribe.py --folder "/path/to/recordings" --out results.jsonl
    python transcribe.py --folder "/path/to/recordings" --out results.jsonl --excel report.xlsx

Resuming:
    Just run the same command again. Files already present in the JSONL
    output (by filename) are skipped automatically. If a run finishes with
    failures, delete the failed lines with --retry-failed and re-run.
"""

import argparse
import json
import os
import sys
import threading
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed

AUDIO_EXTENSIONS = {".mp3": "audio/mpeg", ".wav": "audio/wav", ".m4a": "audio/mp4", ".ogg": "audio/ogg"}
DEEPGRAM_URL = "https://api.deepgram.com/v1/listen"


def parse_args():
    p = argparse.ArgumentParser(description="Fast parallel audio transcription with Deepgram")
    p.add_argument("--folder", required=True, help="Folder containing audio files to transcribe")
    p.add_argument("--out", required=True, help="Output JSONL file (one JSON record per line)")
    p.add_argument("--excel", help="Optional: also write a formatted .xlsx report to this path")
    p.add_argument("--api-key", default=os.environ.get("DEEPGRAM_API_KEY"),
                   help="Deepgram API key (or set DEEPGRAM_API_KEY env var)")
    p.add_argument("--model", default="nova-3", help="Deepgram model (default: nova-3)")
    p.add_argument("--language", default="multi",
                   help="Language code. Use 'multi' for mixed-language audio (e.g. Hindi/English code-switching)")
    p.add_argument("--workers", type=int, default=10, help="Parallel workers (default: 10)")
    p.add_argument("--timeout", type=int, default=300, help="Per-file request timeout in seconds (default: 300)")
    p.add_argument("--retry-failed", action="store_true",
                   help="Remove failed records from an existing --out file before running, so they get retried")
    return p.parse_args()


def load_done(out_path):
    done = set()
    if os.path.exists(out_path):
        with open(out_path, "r", encoding="utf-8") as f:
            for line in f:
                try:
                    rec = json.loads(line)
                    done.add(rec["fileName"])
                except Exception:
                    pass
    return done


def strip_failed(out_path):
    if not os.path.exists(out_path):
        return
    keep = []
    removed = 0
    with open(out_path, "r", encoding="utf-8") as f:
        for line in f:
            rec = json.loads(line)
            if rec.get("error"):
                removed += 1
                continue
            keep.append(line)
    with open(out_path, "w", encoding="utf-8") as f:
        f.writelines(keep)
    print(f"--retry-failed: removed {removed} failed record(s), kept {len(keep)}")


def transcribe_file(path, api_key, model, language, content_type, timeout):
    size = os.path.getsize(path)
    try:
        with open(path, "rb") as f:
            data = f.read()
        url = f"{DEEPGRAM_URL}?model={model}&language={language}&smart_format=true"
        req = urllib.request.Request(
            url,
            data=data,
            headers={"Authorization": f"Token {api_key}", "Content-Type": content_type},
            method="POST",
        )
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            result = json.loads(resp.read().decode("utf-8"))
        transcript = result["results"]["channels"][0]["alternatives"][0]["transcript"]
        duration = result.get("metadata", {}).get("duration", 0)
        return {
            "fileName": os.path.basename(path),
            "fileSize": size,
            "duration": duration,
            "transcript": transcript,
            "error": None,
        }
    except Exception as e:
        return {
            "fileName": os.path.basename(path),
            "fileSize": size,
            "duration": 0,
            "transcript": "",
            "error": str(e),
        }


def write_excel(jsonl_path, excel_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    records = []
    with open(jsonl_path, "r", encoding="utf-8") as f:
        for line in f:
            records.append(json.loads(line))
    records.sort(key=lambda r: r["fileName"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Transcripts"
    headers = ["File Name", "Duration (sec)", "File Size (bytes)", "Transcript", "Error"]
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    for i, r in enumerate(records, start=2):
        ws.cell(row=i, column=1, value=r["fileName"])
        ws.cell(row=i, column=2, value=round(r["duration"], 1) if r["duration"] else 0)
        ws.cell(row=i, column=3, value=r["fileSize"])
        tcell = ws.cell(row=i, column=4, value=r["transcript"] or "")
        tcell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=5, value=r["error"] or "")
        if r["error"]:
            for col in range(1, 6):
                ws.cell(row=i, column=col).fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    widths = [36, 14, 16, 90, 30]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{len(records) + 1}"

    total = len(records)
    success = sum(1 for r in records if not r["error"])
    failed = total - success
    total_min = round(sum(r["duration"] for r in records if r["duration"]) / 60, 1)

    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "Call Recording Transcription Summary"
    ws2["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Total recordings", total),
        ("Successfully transcribed", success),
        ("Failed", failed),
        ("Total audio duration (min)", total_min),
        ("Total audio duration (hours)", round(total_min / 60, 1)),
    ]
    for idx, (label, val) in enumerate(rows, start=3):
        ws2.cell(row=idx, column=1, value=label).font = Font(bold=True)
        ws2.cell(row=idx, column=2, value=val)
    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 20

    wb.save(excel_path)
    print(f"Excel report saved to {excel_path}")


def main():
    args = parse_args()

    if not args.api_key:
        print("ERROR: no API key. Set DEEPGRAM_API_KEY or pass --api-key.", file=sys.stderr)
        sys.exit(1)

    if args.retry_failed:
        strip_failed(args.out)

    if not os.path.isdir(args.folder):
        print(f"ERROR: folder not found: {args.folder}", file=sys.stderr)
        sys.exit(1)

    done = load_done(args.out)
    all_files = []
    for fn in os.listdir(args.folder):
        ext = os.path.splitext(fn)[1].lower()
        if ext not in AUDIO_EXTENSIONS:
            continue
        if fn in done:
            continue
        all_files.append(os.path.join(args.folder, fn))

    print(f"Found {len(all_files) + len(done)} audio files total, {len(done)} already done, {len(all_files)} to process.")
    if not all_files:
        print("Nothing to do.")
    else:
        lock = threading.Lock()
        count = 0
        fail = 0
        print(f"Starting transcription with {args.workers} parallel workers...")
        with ThreadPoolExecutor(max_workers=args.workers) as ex:
            futures = {
                ex.submit(
                    transcribe_file,
                    path,
                    args.api_key,
                    args.model,
                    args.language,
                    AUDIO_EXTENSIONS[os.path.splitext(path)[1].lower()],
                    args.timeout,
                ): path
                for path in all_files
            }
            with open(args.out, "a", encoding="utf-8") as out_f:
                for fut in as_completed(futures):
                    rec = fut.result()
                    if rec["error"]:
                        fail += 1
                    with lock:
                        out_f.write(json.dumps(rec, ensure_ascii=False) + "\n")
                        out_f.flush()
                    count += 1
                    if count % 25 == 0 or count == len(all_files):
                        print(f"progress: {count}/{len(all_files)} (fail={fail})")
        print(f"DONE. Processed={count} fail={fail}")
        if fail:
            print(f"Tip: re-run with --retry-failed to retry the {fail} failed file(s).")

    if args.excel:
        write_excel(args.out, args.excel)


if __name__ == "__main__":
    main()
